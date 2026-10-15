"""packs/schema_induction.py step 2 and experiments/induce.py cluster / name stages (plan §3.3).

A fake embedder (one direction per kind of claim, plus small noise) and a fake naming teacher keep this under a few
seconds; HDBSCAN and PCA are the real ones.
"""

import hashlib
import json
import re

import numpy as np
import pytest
import yaml

from experiments import induce
from experiments.config import load_config
from models.teacher import PROMPTS, Teacher
from packs.schema_induction import (ClusterName, center_by_document, hdbscan_labels, load_claims, naming_sample,
                                    prompt_example_labels, render_name_prompt, summarise_clusters, tune_clusters)

KINDS = ["keeps data for a fixed period", "shares data with service providers", "uses cookies for analytics",
         "tells users how to contact it", "does not knowingly collect data from children"]


class FakeEmbedder:
    """Claims of the same kind share a direction; noise comes from the text hash, so it is deterministic."""

    calls: list[int] = []
    prompts: list[str | None] = []

    def __init__(self, model_id, **kw):
        self.model_id, self.devices, self.normalize, self.truncate_dim, self.dim = model_id, ["cpu"], True, None, 16

    def encode(self, texts, show_progress=False, prompt=None):
        FakeEmbedder.calls.append(len(texts))
        FakeEmbedder.prompts.append(prompt)
        out = []
        for t in texts:
            seed = int(hashlib.sha1(t.encode()).hexdigest()[:8], 16)
            v = np.random.default_rng(seed).normal(0, 0.08, self.dim)
            k = next((i for i, kind in enumerate(KINDS) if kind in t), None)
            if k is None:
                v = np.random.default_rng(seed).normal(0, 1.0, self.dim)  # an outlier
            else:
                v[k] += 1.0
            out.append(v / np.linalg.norm(v))
        return np.asarray(out, dtype=np.float32)


class FakeNamer(Teacher):
    """Names a cluster after the kind its first shown claim belongs to; one kind gets an example label."""

    backend = "fake"

    def __init__(self):
        super().__init__("fake-namer", max_retries=0)
        self.prompts = []

    def _complete(self, conversations, *, json_schema, max_tokens):
        outs = []
        for conv in conversations:
            p = conv[-1]["content"]
            self.prompts.append(p)
            first = re.search(r"^1\. (.*)$", p, re.M).group(1)
            k = next(i for i, kind in enumerate(KINDS) if kind in first)
            name = "policy_change" if k == 0 else "_".join(KINDS[k].split()[:3])
            outs.append(json.dumps({"name": name, "definition": f"States that the company {KINDS[k]}.",
                                    "representative": [3, 1, 1, 99], "single_kind": k != 2,
                                    "split_note": "" if k != 2 else "cookies vs analytics"}))
            self._served()
        return outs


def write_claims(path, genre, n_docs=12, per_kind=4, partial="doc:partial"):
    rows = []
    for d in range(n_docs):
        doc = f"{genre}:{d}"
        j = 0
        for k, kind in enumerate(KINDS):
            for r in range(per_kind if k else per_kind + 1):
                rows.append({"claim_id": f"{doc}#{k}#{j}", "genre": genre, "doc_id": doc, "section_idx": k,
                             "heading": ["Retention", "Sharing", "Cookies", "Contact us", "Children"][k],
                             "claim": f"The company {kind} (variant {d}-{r}).",
                             "type_guess": "loan_period" if (k == 1 and r == 0) else f"kind_{k}", "span": "x"})
                j += 1
        rows.append({"claim_id": f"{doc}#9#0", "genre": genre, "doc_id": doc, "section_idx": 9, "heading": None,
                     "claim": f"Something unrelated number {d}.", "type_guess": "odd", "span": "x"})
    rows += [{"claim_id": f"{partial}#0#{i}", "genre": genre, "doc_id": partial, "section_idx": 0, "heading": None,
              "claim": f"The company {KINDS[0]} (partial {i}).", "type_guess": "kind_0", "span": "x"} for i in range(5)]
    path.write_text("".join(json.dumps(r) + "\n" for r in rows))
    return rows


def test_example_labels_are_read_from_the_extraction_prompt():
    labels = prompt_example_labels("induce_extract")
    assert labels == {"loan_period", "late_fee_policy", "overdue_sanction", "external_context", "policy_change",
                      "evaluation_evidence"}


def test_tune_clusters_picks_a_setting_in_range_with_least_noise():
    rng = np.random.default_rng(0)
    centres = rng.normal(0, 1, (6, 5)) * 6
    Z = np.vstack([c + rng.normal(0, 0.3, (40, 5)) for c in centres] + [rng.normal(0, 8, (10, 5))])
    grid = [{"min_cluster_size": m, "min_samples": None, "selection": "eom"} for m in (5, 30, 300)]
    labels, sweep, best = tune_clusters(Z, grid, target=(4, 8), max_cluster_share=0.5)
    assert len(sweep) == 3 and {"n_clusters", "noise_rate", "largest_share", "eligible"} <= sweep[0].keys()
    assert sweep[2]["n_clusters"] < 4 and not sweep[2]["eligible"]
    assert 4 <= sweep[best]["n_clusters"] <= 8 and len(set(labels) - {-1}) == sweep[best]["n_clusters"]
    eligible = [r for r in sweep if r["eligible"]]
    assert sweep[best]["noise_rate"] == min(r["noise_rate"] for r in eligible)


def test_centering_by_document_removes_a_shared_topic_direction():
    """Each document has a strong topic direction shared by all its claims; the kind of claim is weaker. Raw vectors
    cluster by document (topic); centred vectors cluster by kind (the genre move)."""
    rng = np.random.default_rng(0)
    n_docs, n_kinds, dim = 6, 4, 24
    topic = rng.normal(0, 1, (n_docs, dim)) * 3
    kind = np.eye(dim)[:n_kinds] * 1.5
    X, docs, kinds = [], [], []
    for d in range(n_docs):
        for k in range(n_kinds):
            for _ in range(8):
                X.append(topic[d] + kind[k] + rng.normal(0, 0.1, dim))
                docs.append(f"d{d}")
                kinds.append(k)
    X = np.asarray(X)
    X /= np.linalg.norm(X, axis=1, keepdims=True)
    kinds = np.asarray(kinds)

    def purity(labels, truth):
        ok = labels >= 0
        return sum(np.bincount(truth[ok & (labels == c)]).max() for c in set(labels[ok])) / max(1, ok.sum())

    raw = hdbscan_labels(X, min_cluster_size=10, min_samples=5, selection="eom")
    Y = center_by_document(X, docs, alpha=1.0)
    centred = hdbscan_labels(Y, min_cluster_size=10, min_samples=5, selection="eom")
    assert purity(raw, kinds) < 0.5 and purity(centred, kinds) > 0.95
    assert np.allclose(np.linalg.norm(Y, axis=1), 1.0, atol=1e-5)
    assert np.allclose(center_by_document(X, docs, alpha=0.0), X, atol=1e-6)


def test_summaries_order_members_by_centrality_and_describe_placement(tmp_path):
    rows = write_claims(tmp_path / "c.jsonl", "privacy", n_docs=3)
    X = FakeEmbedder("x").encode([r["claim"] for r in rows])
    labels = np.array([next((i for i, k in enumerate(KINDS) if k in r["claim"]), -1) for r in rows])
    s = summarise_clusters(rows, labels, X)
    assert [c["size"] for c in s] == sorted([c["size"] for c in s], reverse=True)
    assert s[0]["id"] == "c000" and s[0]["heading_top"][0][0] == "retention" and s[0]["documents"] == 4
    assert all(0.0 <= c["position_mean"] <= 1.0 for c in s)
    sims = [float(X[[r["claim_id"] for r in rows].index(m)] @ np.asarray(s[1]["centroid"])) for m in s[1]["members"]]
    assert sims == sorted(sims, reverse=True)
    shown = naming_sample(s[1]["members"], n=6, seed=0)
    assert len(shown) == 6 and shown[:3] == s[1]["members"][:3] and len(set(shown)) == 6
    assert naming_sample(s[1]["members"], n=6, seed=0) == shown


def test_name_prompt_renders_for_every_genre():
    cfg = load_config("configs/induce.yaml")
    for genre, info in cfg["genre_info"].items():
        p = render_name_prompt(["The company keeps data.", "The company deletes data."], info, cluster_size=40,
                               n_documents=12, type_hints=[("retention_period", 30), ("deletion", 4)])
        assert info["document_kind"] in p and "1. The company keeps data." in p and "retention_period (30)" in p
        assert "{{" not in p


@pytest.fixture
def setup(tmp_path, monkeypatch):
    for genre in ("privacy", "papers"):
        write_claims(tmp_path / f"claims_{genre}.jsonl", genre)
    cfg = load_config("configs/induce.yaml")
    cfg["paths"].update(embeddings=str(tmp_path / "emb"), packs=str(tmp_path / "packs"))
    cfg["run"]["seed"] = 0
    c = cfg["induce"]["cluster"]
    c["claims"] = {g: str(tmp_path / f"claims_{g}.jsonl") for g in ("privacy", "papers")}
    c["exclude_docs"] = ["doc:partial"]
    c["reduce"] = {"method": "pca", "dims": 8}
    c["grid"] = {"min_cluster_size": [3, 8, 200], "min_samples": [None], "selection": ["eom"]}
    cfg["schema"]["target_clusters"] = [4, 8]
    monkeypatch.setattr("engine.embed.Embedder", FakeEmbedder)
    FakeEmbedder.calls, FakeEmbedder.prompts = [], []
    return cfg, tmp_path


def test_cluster_stage_end_to_end_and_embedding_reuse(setup):
    cfg, tmp = setup
    cfg["induce"]["stage"] = "cluster"
    (tmp / "run1").mkdir()
    m = induce.main(cfg, tmp / "run1")
    assert m["privacy/claims"] == 12 * 22 and m["privacy/excluded_claims"] == 5 and m["privacy/documents"] == 12
    assert m["privacy/n_clusters"] == 5 and 0 < m["privacy/noise_rate"] < 0.1
    assert len(m["privacy/sweep"]) == 3 and m["privacy/chosen"]["min_cluster_size"] in (3, 8)
    clusters = json.loads((tmp / "run1" / "clusters_privacy.json").read_text())
    assert len(clusters) == 5 and "centroid" not in clusters[0]
    assert np.load(tmp / "run1" / "centroids_privacy.npy").shape == (5, 16)
    assign = [json.loads(l) for l in open(tmp / "run1" / "assignments_privacy.jsonl")]
    assert len(assign) == 12 * 22 and {a["cluster"] for a in assign} == {"noise"} | {c["id"] for c in clusters}
    assert sum(FakeEmbedder.calls) == 2 * 12 * 22
    instruction = PROMPTS.get("induce_cluster_embed")
    assert m["embed_prompt"] == f"induce_cluster_embed {instruction.tag}"
    assert FakeEmbedder.prompts[0].startswith("Instruct: ") and FakeEmbedder.prompts[0].endswith("Query: ")
    assert "website privacy policy" in FakeEmbedder.prompts[0]  # genre_info.privacy.document_kind
    stores = sorted(p.name for p in (tmp / "emb").glob("*/privacy/*"))
    assert stores == [f"claims__induce_cluster_embed_v{instruction.version}_{instruction.hash}"]

    FakeEmbedder.calls = []
    cfg["schema"]["hdbscan"].update(min_cluster_size=m["privacy/chosen"]["min_cluster_size"])
    (tmp / "run2").mkdir()
    m2 = induce.main(cfg, tmp / "run2")
    assert sum(FakeEmbedder.calls) == 0 and m2["privacy/embedded_now"] == 0  # every vector came from the store
    assert len(m2["privacy/sweep"]) == 1 and m2["privacy/n_clusters"] == 5
    assert json.loads((tmp / "run2" / "clusters_privacy.json").read_text()) == clusters


def test_name_stage_writes_draft_and_flags_leakage_and_duplicates(setup, monkeypatch):
    cfg, tmp = setup
    cfg["induce"]["stage"] = "cluster"
    (tmp / "c").mkdir()
    (tmp / "c" / "config.yaml").write_text(yaml.safe_dump(cfg))
    induce.main(cfg, tmp / "c")
    namer = FakeNamer()
    monkeypatch.setattr("models.teacher.load_teacher", lambda *a, **k: namer)
    PROMPTS.reset_usage()
    cfg["induce"]["stage"] = "name"
    cfg["induce"]["name"].update(from_run=str(tmp / "c"), claims_per_cluster=6, write_pack=True, limit=None)
    (tmp / "n").mkdir()
    m = induce.main(cfg, tmp / "n")
    assert PROMPTS.usage()["induce_name"] == PROMPTS.get("induce_name").tag
    assert m["privacy/named"] == 5 and m["privacy/failed"] == 0 and m["privacy/mixed"] == 1
    assert m["privacy/example_label_names"] == ["policy_change"]
    draft = yaml.safe_load((tmp / "packs" / "privacy" / "schema_draft.yaml").read_text())
    assert draft["status"] == "draft" and draft["induced"]["clusters"] == 5
    assert draft["induced"]["naming_prompt"] == f"induce_name {PROMPTS.get('induce_name').tag}"
    t0 = draft["types"][0]
    assert {"id", "name", "definition", "size", "documents", "representative_claims", "single_kind", "flags",
            "type_guess_top", "heading_top", "position_mean"} <= t0.keys()
    assert len(t0["representative_claims"]) == 3  # out-of-range and repeated picks replaced by central claims
    assert len({c["claim_id"] for c in t0["representative_claims"]}) == 3
    leak = [t for t in draft["types"] if "example_label_name" in t["flags"]]
    assert [t["name"] for t in leak] == ["policy_change"]
    sharing = next(t for t in draft["types"] if t["name"] == "shares_data_with")
    assert sharing["example_label_share"] > 0  # members whose type_guess was an example label

    # a limited (iteration) run never touches the pack, and duplicate names are flagged as merge candidates
    (tmp / "packs" / "privacy" / "schema_draft.yaml").unlink()
    monkeypatch.setattr(FakeNamer, "_complete", lambda self, convs, **k: [json.dumps(
        {"name": "same_name", "definition": "d", "representative": [1, 2, 3], "single_kind": True, "split_note": ""})
        for _ in convs])
    cfg["induce"]["name"]["limit"] = 2
    (tmp / "n2").mkdir()
    m2 = induce.main(cfg, tmp / "n2")
    assert m2["privacy/named"] == 2 and m2["privacy/duplicate_names"] == 1
    assert not (tmp / "packs" / "privacy" / "schema_draft.yaml").exists()
    d2 = yaml.safe_load((tmp / "n2" / "schema_draft_privacy.yaml").read_text())
    assert all("duplicate_name" in t["flags"] for t in d2["types"])


def test_review_stage_writes_the_workbook(setup, monkeypatch):
    from openpyxl import load_workbook

    cfg, tmp = setup
    cfg["induce"]["stage"] = "cluster"
    (tmp / "c").mkdir()
    (tmp / "c" / "config.yaml").write_text(yaml.safe_dump(cfg))
    induce.main(cfg, tmp / "c")
    monkeypatch.setattr("models.teacher.load_teacher", lambda *a, **k: FakeNamer())
    cfg["induce"]["stage"] = "name"
    cfg["induce"]["name"].update(from_run=str(tmp / "c"), claims_per_cluster=6, write_pack=False)
    (tmp / "n").mkdir()
    induce.main(cfg, tmp / "n")
    (tmp / "seeds.yaml").write_text(yaml.safe_dump({"items": [{"id": "s1", "title": "retention"}]}))
    clusters = json.loads((tmp / "c" / "clusters_privacy.json").read_text())
    (tmp / "proposals_privacy.yaml").write_text(yaml.safe_dump({"types": [
        {"id": "retention", "name": "retention", "definition": "States how long data is kept.",
         "clusters": [clusters[0]["id"]], "positive_examples": clusters[0]["members"][:2],
         "near_miss": clusters[1]["members"][0], "section_hint": "middle", "seeds": ["s1"]}]}))
    cfg["induce"]["stage"] = "review"
    cfg["induce"]["review"].update(cluster_run=str(tmp / "c"), name_run=str(tmp / "n"),
                                   proposals=str(tmp / "proposals_{genre}.yaml"), seeds=str(tmp / "seeds.yaml"),
                                   out_dir=str(tmp / "out"), noise_sample=5)
    cfg["pack"] = "privacy"
    (tmp / "r").mkdir()
    m = induce.main(cfg, tmp / "r")
    assert m["privacy/clusters"] == 5 and m["privacy/proposed_types"] == 1 and m["privacy/noise_sampled"] == 5
    wb = load_workbook(tmp / "out" / "schema_review_privacy.xlsx")
    assert wb["Types"]["A2"].value == "retention"
    assert wb["Types"]["D2"].value.startswith("The company keeps data")  # claim ids resolved to text
    assert wb["Noise sample"].max_row == 6 and (tmp / "r" / "schema_review_privacy.xlsx").exists()

    # the reviewer's pass: accept everything, prune the clusters without a proposal, map the seed
    ws = wb["Clusters"]
    h = [c.value for c in ws[1]]
    for r in range(2, ws.max_row + 1):
        proposed = ws.cell(r, h.index("Proposed action") + 1).value
        ws.cell(r, h.index("Your action") + 1).value = "accept" if proposed != "review" else "prune"
    wb["Seeds"]["E2"] = "accept"
    wb.save(tmp / "out" / "schema_review_privacy.xlsx")
    cfg["induce"]["stage"] = "candidate"
    cfg["induce"]["candidate"].update(workbook=str(tmp / "out" / "schema_review_{genre}.xlsx"), cluster_run=str(tmp / "c"))
    cfg["schema"]["min_types"] = 1
    (tmp / "k").mkdir()
    m = induce.main(cfg, tmp / "k")
    cand = yaml.safe_load((tmp / "packs" / "privacy" / "schema_candidate.yaml").read_text())
    assert m["privacy/types"] == 1 and m["privacy/pruned_clusters"] == 4 and cand["status"] == "candidate"
    assert cand["types"][0]["source_clusters"] == [clusters[0]["id"]]
    assert any("label_source" in p for p in m["privacy/gate_problems"])  # privacy needs OPP-115 label sources
