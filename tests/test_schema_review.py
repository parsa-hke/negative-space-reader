"""packs/schema_review.py: the 3.4 review workbook (build, read back) and schema validation for the 3.6 gate."""

import numpy as np
import pytest
from openpyxl import load_workbook

from packs.schema_review import (build_workbook, nearest_clusters, read_workbook, schema_from_review,
                                 validate_schema)


def claims_and_clusters():
    claims = {f"d{d}#{k}#{j}": {"claim_id": f"d{d}#{k}#{j}", "doc_id": f"d{d}", "section_idx": k,
                                "heading": "H", "claim": f"The company does kind {k} ({d}.{j}).",
                                "type_guess": f"kind_{k}"} for d in range(5) for k in range(4) for j in range(3)}
    clusters = [{"id": f"c00{k}", "size": 15, "documents": 5, "cohesion": 0.9, "position_mean": 0.5,
                 "type_guess_top": [[f"kind_{k}", 15]], "heading_top": [["h", 15]],
                 "members": [c for c in claims if f"#{k}#" in c]} for k in range(4)]
    names = {f"c00{k}": {"name": f"kind_{k}_statement", "definition": f"States kind {k}.", "flags": [],
                         "split_note": "", "representative_claims": [{"claim_id": f"d0#{k}#0",
                                                                      "claim": claims[f"d0#{k}#0"]["claim"]}]}
             for k in range(4)}
    centroids = np.eye(4, 8, dtype=np.float32) + 0.1
    noise = [{"claim_id": "d9#0#0", "doc_id": "d9", "heading": None, "claim": "Something odd.",
              "nearest": "c000", "similarity": 0.4}]
    return claims, clusters, names, centroids, noise


PROPOSALS = {
    "types": [
        {"id": "kind_a", "name": "kind_a_statement", "definition": "States kind 0 or 1.", "clusters": ["c000", "c001"],
         "positive_examples": ["d0#0#0", "d1#1#0"], "near_miss": "d0#2#0", "severity_weight": 1.5,
         "section_hint": "middle", "label_source": ["Data Retention"], "seeds": ["s1"]},
        {"id": "kind_c", "name": "kind_c_statement", "definition": "States kind 2.", "clusters": ["c002"],
         "positive_examples": ["d0#2#0", "d1#2#0"], "near_miss": "d0#0#0", "section_hint": "end",
         "label_source": ["teacher"]}],
    "clusters": {"c003": {"action": "prune", "reason": "boilerplate"}},
    "seeds": {"s2": {"exclude": "not a practice"}},
}
SEEDS = {"items": [{"id": "s1", "title": "retention"}, {"id": "s2", "title": "other"}]}


def test_nearest_clusters():
    c = np.array([[1, 0], [0.9, 0.1], [0, 1]], dtype=np.float32)
    c /= np.linalg.norm(c, axis=1, keepdims=True)
    nn = nearest_clusters(c, ["a", "b", "c"], k=1)
    assert nn["a"][0][0] == "b" and nn["c"][0][0] in ("a", "b")


def test_workbook_round_trip_accepting_proposals(tmp_path):
    claims, clusters, names, centroids, noise = claims_and_clusters()
    path = tmp_path / "review.xlsx"
    build_workbook(path, genre="privacy", clusters=clusters, names=names, centroids=centroids, claims=claims,
                   noise=noise, proposals=PROPOSALS, seeds=SEEDS, n_documents=5, require_label_source=True)
    wb = load_workbook(path)
    assert {"Read me", "Summary", "Clusters", "Types", "Seeds", "Claims", "Noise sample"} <= set(wb.sheetnames)
    ws = wb["Clusters"]
    header = [c.value for c in ws[1]]
    for col in ("Cluster", "Proposed action", "Proposed type", "Your action", "Your type", "Notes"):
        assert col in header
    rows = {r[0]: dict(zip(header, r)) for r in ws.iter_rows(min_row=2, values_only=True) if r[0]}
    assert rows["c000"]["Proposed action"] == "merge" and rows["c000"]["Proposed type"] == "kind_a"
    assert rows["c002"]["Proposed action"] == "keep" and rows["c003"]["Proposed action"] == "prune"
    assert rows["c000"]["Your action"] is None  # the human decides; nothing is pre-accepted

    # undecided clusters block the schema
    review = read_workbook(path)
    with pytest.raises(ValueError, match="undecided"):
        schema_from_review(review, claims, genre="privacy")

    # the human accepts every proposal, then overrides one: c003 is kept as its own new type
    ycol = header.index("Your action") + 1
    tcol = header.index("Your type") + 1
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 1).value:
            ws.cell(r, ycol).value = "accept"
    r3 = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value == "c003")
    ws.cell(r3, ycol).value, ws.cell(r3, tcol).value = "keep", "kind_d"
    t = wb["Types"]
    th = [c.value for c in t[1]]
    t.append([{"Type id": "kind_d", "Name": "kind_d_statement", "Definition": "States kind 3.",
               "Positive example 1": "The company does kind 3 (0.0).", "Positive example 2": "The company does kind 3 (1.0).",
               "Near miss": "The company does kind 2 (0.0).", "Severity weight": 1.0, "Section hint": "end",
               "Label source": "Other :: Other Type = Practice not covered"}.get(h) for h in th])
    s = wb["Seeds"]
    sh = [c.value for c in s[1]]
    for r in range(2, s.max_row + 1):
        s.cell(r, sh.index("Your mapping") + 1).value = "accept"
    wb.save(path)

    schema = schema_from_review(read_workbook(path), claims, genre="privacy")
    ids = [t["id"] for t in schema["types"]]
    assert ids == ["kind_a", "kind_c", "kind_d"]
    a = schema["types"][0]
    assert a["source_clusters"] == ["c000", "c001"] and a["severity_weight"] == 1.5
    assert a["positive_examples"] == [claims["d0#0#0"]["claim"], claims["d1#1#0"]["claim"]]
    assert a["near_miss"] == claims["d0#2#0"]["claim"] and a["label_source"] == ["Data Retention"]
    assert schema["types"][2]["source_clusters"] == ["c003"]
    assert schema["seed_mapping"]["s1"] == {"types": ["kind_a"]} and schema["seed_mapping"]["s2"] == {"exclude": "not a practice"}

    problems = validate_schema(schema, min_types=3, max_types=60, require_label_source=True)
    assert problems == []
    assert any("40" in p for p in validate_schema(schema, min_types=40, max_types=60, require_label_source=True))


def test_validate_schema_lists_every_missing_field():
    schema = {"types": [{"id": "x", "name": "x", "definition": "", "positive_examples": ["a"], "near_miss": "",
                         "severity_weight": 1.0, "section_hint": "", "label_source": []},
                        {"id": "x", "name": "y", "definition": "d", "positive_examples": ["a", "b"], "near_miss": "n",
                         "severity_weight": 1.0, "section_hint": "s", "label_source": ["teacher"]}],
              "seed_mapping": {"s1": {}}}
    p = validate_schema(schema, min_types=1, max_types=5, require_label_source=True)
    joined = "\n".join(p)
    for needle in ("duplicate type id x", "x: definition", "x: 2 positive examples", "x: near_miss",
                   "x: section_hint", "x: label_source", "seed s1"):
        assert needle in joined


def test_complete_proposals_fills_examples_near_miss_and_section_hint():
    from packs.schema_review import complete_proposals

    claims, clusters, names, centroids, _ = claims_and_clusters()
    clusters[0]["heading_top"], clusters[0]["position_mean"] = [["(no heading)", 9], ["data retention", 5]], 0.9
    clusters[1]["heading_top"] = [["(no heading)", 15]]  # size-weighted position (0.9, 0.5) -> 0.7
    props = {"types": [{"id": "a", "clusters": ["c000", "c001"]},
                       {"id": "b", "clusters": [], "positive_examples": ["d0#3#0", "d1#3#0"], "near_miss": "d0#2#0",
                        "section_hint": "given"}]}
    out = complete_proposals(props, clusters, names, centroids)
    a, b = out["types"]
    assert a["positive_examples"] == ["d0#0#0", "d0#1#0"]  # first representative of each source cluster
    assert a["near_miss"] in {"d0#2#0", "d0#3#0"}  # from the nearest cluster outside the type
    assert "data retention" in a["section_hint"] and "second half" in a["section_hint"]
    assert set(a["auto"]) == {"positive_examples", "near_miss", "section_hint"}
    assert b["positive_examples"] == ["d0#3#0", "d1#3#0"] and b["section_hint"] == "given" and not b.get("auto")
