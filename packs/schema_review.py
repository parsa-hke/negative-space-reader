"""Schema review (plan §3.4) and the checks of the schema gate (§3.6). Genre-agnostic: a pack's seeds and the
proposals are data.

    build_workbook(...)        cluster run + names + proposals + seeds -> an .xlsx for the human pass: one row per
                               cluster (statistics, teacher name, representative claims, nearest clusters, the
                               proposed action and type), the proposed types with every required field, the seed
                               cross-check, claims per cluster, a sample of noise claims; yellow columns are the
                               reviewer's ("Your action" etc.), nothing is pre-accepted
    read_workbook(path)        the reviewer's workbook -> plain rows
    schema_from_review(...)    rows -> schema dict (refuses undecided clusters, unknown type ids, unmapped seeds)
    validate_schema(...)       the §3.6 checks: 40-60 types, unique ids, definition, 2 positives, 1 near-miss,
                               section hint, severity weight, label source where the pack needs one, every seed
                               mapped or excluded with a reason

Proposals (written by Claude after reading the named clusters, `reports/schema_review/proposals_<genre>.yaml`):
    types:    [{id, name, definition, clusters: [cNNN], positive_examples: [claim id or text] x2,
                near_miss: claim id or text, severity_weight, section_hint, label_source: [..], seeds: [seed ids]}]
    clusters: {cNNN: {action: prune | split, reason}}       (clusters not listed in a type and not here: "review")
    seeds:    {seed id: {exclude: reason}}
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

ACTIONS = ("accept", "keep", "merge", "split", "prune")
REQUIRED = ("definition", "near_miss", "section_hint")
TYPE_COLUMNS = ["Type id", "Name", "Definition", "Positive example 1", "Positive example 2", "Near miss",
                "Severity weight", "Section hint", "Label source", "Seeds", "Source clusters (proposed)",
                "Claims (your decisions)", "Notes"]
CLUSTER_COLUMNS = ["Cluster", "Size", "Documents", "Document share", "Cohesion", "Position (0 first, 1 last)",
                   "Top headings", "Teacher name", "Teacher definition", "Flags", "Split note",
                   "Representative 1", "Representative 2", "Representative 3", "Top extraction labels",
                   "Nearest clusters (cosine)", "Proposed action", "Proposed type", "Reason",
                   "Your action", "Your type", "Notes", "Effective type"]
SEED_COLUMNS = ["Seed", "Title", "Proposed types", "Proposed exclusion", "Your mapping", "Notes"]


# ------------------------------------------------------------------------------------ helpers
def nearest_clusters(centroids: "np.ndarray", ids: list[str], *, k: int = 3) -> dict[str, list[tuple[str, float]]]:
    """For each cluster, the k clusters with the most similar centroid (cosine): merge candidates."""
    import numpy as np

    c = centroids / np.maximum(np.linalg.norm(centroids, axis=1, keepdims=True), 1e-12)
    sims = c @ c.T
    np.fill_diagonal(sims, -np.inf)
    out = {}
    for i, cid in enumerate(ids):
        order = np.argsort(-sims[i])[:k]
        out[cid] = [(ids[j], round(float(sims[i, j]), 3)) for j in order]
    return out


def _text(ref: Any, claims: dict[str, dict[str, Any]]) -> str:
    """A proposal's example is a claim id (resolved to its text) or literal text."""
    if ref is None:
        return ""
    return claims[ref]["claim"] if ref in claims else str(ref)


def _split_list(value: Any) -> list[str]:
    if value is None:
        return []
    if isinstance(value, (list, tuple)):
        return [str(v).strip() for v in value if str(v).strip()]
    parts = str(value).replace(";", "\n").replace(",", "\n").split("\n")
    return [p.strip() for p in parts if p.strip()]


def _label_list(value: Any) -> list[str]:
    """Label sources contain commas inside values, so only newlines and semicolons separate entries."""
    if value is None:
        return []
    if isinstance(value, (list, tuple)):
        return [str(v).strip() for v in value if str(v).strip()]
    return [p.strip() for p in str(value).replace(";", "\n").split("\n") if p.strip()]


def proposed_actions(cluster_ids: list[str], proposals: dict[str, Any]) -> dict[str, tuple[str, str, str]]:
    """cluster id -> (action, type ids, reason) from the proposals."""
    out: dict[str, tuple[str, str, str]] = {}
    owners: dict[str, list[str]] = {}
    for t in proposals.get("types", []):
        for c in t.get("clusters", []):
            owners.setdefault(c, []).append(t["id"])
    for cid in cluster_ids:
        explicit = (proposals.get("clusters") or {}).get(cid)
        if explicit:
            out[cid] = (explicit["action"], ", ".join(owners.get(cid, [])), explicit.get("reason", ""))
        elif cid in owners:
            t = owners[cid]
            merged = any(len(tt.get("clusters", [])) > 1 for tt in proposals["types"] if tt["id"] in t)
            out[cid] = ("split" if len(t) > 1 else "merge" if merged else "keep", ", ".join(t), "")
        else:
            out[cid] = ("review", "", "no proposal")
    return out


# ------------------------------------------------------------------------------------ build
def build_workbook(path: str | Path, *, genre: str, clusters: list[dict[str, Any]], names: dict[str, dict[str, Any]],
                   centroids: "np.ndarray", claims: dict[str, dict[str, Any]], noise: list[dict[str, Any]],
                   proposals: dict[str, Any], seeds: dict[str, Any], n_documents: int, require_label_source: bool,
                   claims_per_cluster: int = 20, provenance: dict[str, Any] | None = None) -> None:
    import random

    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    font, bold = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
    yellow = PatternFill("solid", start_color="FFFF00")
    grey = PatternFill("solid", start_color="D9D9D9")
    wrap = Alignment(wrap_text=True, vertical="top")

    def sheet(wb: Any, title: str, columns: list[str], widths: dict[str, int], editable: set[str]) -> Any:
        ws = wb.create_sheet(title)
        ws.append(columns)
        for i, col in enumerate(columns, 1):
            cell = ws.cell(1, i)
            cell.font, cell.alignment = bold, wrap
            cell.fill = yellow if col in editable else grey
            ws.column_dimensions[get_column_letter(i)].width = widths.get(col, 14)
        ws.freeze_panes = "B2"
        return ws

    def style_rows(ws: Any, columns: list[str], editable: set[str]) -> None:
        for row in ws.iter_rows(min_row=2):
            for cell, col in zip(row, columns):
                cell.font, cell.alignment = font, wrap
                if col in editable:
                    cell.fill = yellow

    wb = Workbook()
    wb.remove(wb.active)
    ids = [c["id"] for c in clusters]
    nn = nearest_clusters(centroids, ids)
    proposed = proposed_actions(ids, proposals)

    # -- Read me
    ws = wb.create_sheet("Read me")
    lines = [
        f"Schema review workbook, genre: {genre} (plan task 3.4: merge, split and prune clusters to 40-60 types)",
        "",
        "Yellow columns are yours. Grey columns are computed (Clusters, Claims, Noise sample) or proposed by Claude.",
        "Clusters sheet, 'Your action' (dropdown): accept = take the proposed action and type as they are;",
        "  keep = the cluster is one type (give it in 'Your type'); merge = it joins the type in 'Your type' with",
        "  other clusters; split = it holds several types (list them in 'Your type', comma-separated);",
        "  prune = not a genre type (say why in 'Notes'). A row left blank blocks the schema.",
        "Types sheet: the proposed types with every field the frozen schema needs. Edit any cell, delete rows,",
        "  add rows. Type ids must match what the Clusters sheet points to. Examples are claim texts.",
        "  Severity weight: 1.0 by default; 1.5 for privacy types that GDPR Art. 13/14 or CCPA require.",
        "  Section hint: where peers usually state it (used by the ghost-paragraph view).",
        "  Label source (privacy): OPP-115 values that count as a positive, one per line, written",
        "  'Category :: Attribute = Value' or just 'Category' (lines are alternatives; '&' joins attribute values that",
        "  must all hold); 'teacher' when no OPP-115 value maps (the type then gets teacher labels in Phase 4).",
        "Notes column 'auto-filled, check: ...' = Claude picked that field automatically from the clusters; a human",
        "  should look at it. Flag 'example_label_name' = the name contains a label from the extraction prompt's",
        "  examples (possible leakage); 'mixed' = the naming teacher saw two kinds of statement in the cluster.",
        "Seeds sheet, 'Your mapping': accept, or type ids (comma-separated), or 'exclude: <reason>'.",
        "Summary sheet: counts update as you edit (open in Excel, Numbers or LibreOffice).",
        "Example Clusters row: c012 | ... | Your action = merge | Your type = third_party_sharing | Notes = same move as c000",
        "",
        "Provenance: " + ", ".join(f"{k}={v}" for k, v in (provenance or {}).items()),
    ]
    for line in lines:
        ws.append([line])
    ws.column_dimensions["A"].width = 120
    for row in ws.iter_rows():
        row[0].font = font
    ws["A1"].font = bold

    # -- Clusters
    edit_c = {"Your action", "Your type", "Notes"}
    wc = sheet(wb, "Clusters", CLUSTER_COLUMNS, {"Teacher definition": 40, "Representative 1": 40,
                                                   "Representative 2": 40, "Representative 3": 40, "Top headings": 24,
                                                   "Top extraction labels": 28, "Nearest clusters (cosine)": 24,
                                                   "Teacher name": 24, "Reason": 28, "Notes": 30, "Split note": 20,
                                                   "Proposed type": 22, "Your type": 22, "Effective type": 22}, edit_c)
    col = {c: i + 1 for i, c in enumerate(CLUSTER_COLUMNS)}
    L = {c: get_column_letter(i) for c, i in col.items()}
    for r, c in enumerate(clusters, 2):
        nm = names.get(c["id"], {})
        reps = [x["claim"] for x in nm.get("representative_claims", [])] + ["", "", ""]
        action, types, reason = proposed[c["id"]]
        wc.append([c["id"], c["size"], c["documents"], round(c["documents"] / max(1, n_documents), 3), c["cohesion"],
                   c["position_mean"], "; ".join(f"{h} ({n})" for h, n in c["heading_top"]), nm.get("name", ""),
                   nm.get("definition", ""), ", ".join(nm.get("flags", [])), nm.get("split_note", ""),
                   reps[0], reps[1], reps[2], ", ".join(f"{t} ({n})" for t, n in c["type_guess_top"]),
                   ", ".join(f"{i} {s:.2f}" for i, s in nn[c["id"]]), action, types, reason, None, None, None,
                   f'=IF({L["Your action"]}{r}="accept",{L["Proposed type"]}{r},{L["Your type"]}{r})'])
    style_rows(wc, CLUSTER_COLUMNS, edit_c)
    dv = DataValidation(type="list", formula1='"' + ",".join(ACTIONS) + '"', allow_blank=True)
    wc.add_data_validation(dv)
    dv.add(f'{L["Your action"]}2:{L["Your action"]}{len(clusters) + 1}')
    wc.cell(1, col["Document share"]).comment = Comment(f"documents / {n_documents} sampled documents", "Claude")
    wc.cell(1, col["Cohesion"]).comment = Comment("mean cosine of members to the centroid (clustering space)", "Claude")
    wc.auto_filter.ref = f"A1:{L['Effective type']}{len(clusters) + 1}"

    # -- Types
    edit_t = set(TYPE_COLUMNS) - {"Claims (your decisions)"}
    wt = sheet(wb, "Types", TYPE_COLUMNS, {"Definition": 50, "Positive example 1": 45, "Positive example 2": 45,
                                            "Near miss": 45, "Label source": 36, "Section hint": 24, "Name": 26,
                                            "Type id": 24, "Notes": 30, "Source clusters (proposed)": 20}, edit_t)
    last = len(clusters) + 1
    for r, t in enumerate(proposals.get("types", []), 2):
        pos = [_text(x, claims) for x in t.get("positive_examples", [])] + ["", ""]
        wt.append([t["id"], t.get("name", t["id"]), t.get("definition", ""), pos[0], pos[1],
                   _text(t.get("near_miss"), claims), t.get("severity_weight", 1.0), t.get("section_hint", ""),
                   "\n".join(t.get("label_source", [])), ", ".join(t.get("seeds", [])), ", ".join(t.get("clusters", [])),
                   f'=SUMIF(Clusters!${L["Effective type"]}$2:${L["Effective type"]}${last},A{r},'
                   f'Clusters!$B$2:$B${last})',
                   "; ".join(x for x in (t.get("notes", ""), ("auto-filled, check: " + ", ".join(t["auto"]))
                                         if t.get("auto") else "") if x)])
    style_rows(wt, TYPE_COLUMNS, edit_t)

    # -- Seeds
    edit_s = {"Your mapping", "Notes"}
    wsd = sheet(wb, "Seeds", SEED_COLUMNS, {"Title": 50, "Proposed types": 30, "Proposed exclusion": 40,
                                            "Your mapping": 30, "Notes": 30}, edit_s)
    seed_types: dict[str, list[str]] = {}
    for t in proposals.get("types", []):
        for s in t.get("seeds", []):
            seed_types.setdefault(s, []).append(t["id"])
    for item in seeds.get("items", []):
        excl = ((proposals.get("seeds") or {}).get(item["id"]) or {}).get("exclude", "")
        title = item.get("title") or item.get("category", "")
        wsd.append([item["id"], title, ", ".join(seed_types.get(item["id"], [])), excl, None, None])
    style_rows(wsd, SEED_COLUMNS, edit_s)

    # -- Claims (most typical half, then a seeded random sample of the rest)
    wl = sheet(wb, "Claims", ["Cluster", "Rank", "Kind", "Claim", "Document", "Heading"],
               {"Claim": 90, "Document": 26, "Heading": 30}, set())
    rng = random.Random(0)
    for c in clusters:
        head = c["members"][: claims_per_cluster // 2]
        rest = c["members"][claims_per_cluster // 2:]
        sample = rng.sample(rest, min(claims_per_cluster - len(head), len(rest)))
        for i, m in enumerate(head + sample, 1):
            x = claims[m]
            wl.append([c["id"], i, "typical" if i <= len(head) else "random", x["claim"], x["doc_id"], x.get("heading")])
    style_rows(wl, ["Cluster", "Rank", "Kind", "Claim", "Document", "Heading"], set())

    # -- Noise sample
    wn = sheet(wb, "Noise sample", ["Claim", "Document", "Heading", "Nearest cluster", "Similarity"],
               {"Claim": 90, "Document": 26, "Heading": 30}, set())
    for x in noise:
        wn.append([x["claim"], x["doc_id"], x.get("heading"), x["nearest"], x["similarity"]])
    style_rows(wn, ["Claim", "Document", "Heading", "Nearest cluster", "Similarity"], set())

    # -- Summary (formulas, so it follows the reviewer's edits)
    wsum = wb.create_sheet("Summary", 1)
    ya, yt = L["Your action"], L["Your type"]
    s_last = len(seeds.get("items", [])) + 1
    rows = [
        ("Clusters", f"=COUNTA(Clusters!A2:A{last})"),
        ("Clusters still undecided", f'=SUMPRODUCT((Clusters!A2:A{last}<>"")*(Clusters!{ya}2:{ya}{last}=""))'),
        ("Clusters pruned (your decisions)", f'=COUNTIF(Clusters!{ya}2:{ya}{last},"prune")'),
        ("Claims in pruned clusters", f'=SUMIF(Clusters!{ya}2:{ya}{last},"prune",Clusters!B2:B{last})'),
        ("Types on the Types sheet", "=COUNTA(Types!A2:A500)"),
        ("Types within 40-60?", '=IF(AND(B6>=40,B6<=60),"yes","no")'),
        ("Types missing a required field",
         '=SUMPRODUCT((Types!A2:A500<>"")*(((Types!C2:C500="")+(Types!D2:D500="")+(Types!E2:E500="")'
         '+(Types!F2:F500="")+(Types!H2:H500="")' + ('+(Types!I2:I500="")' if require_label_source else "") + ')>0))'),
        ("Seeds still undecided", f'=COUNTBLANK(Seeds!E2:E{s_last})'),
        ("Kept-or-merged clusters with no type given",
         f'=SUMPRODUCT((Clusters!{ya}2:{ya}{last}<>"")*(Clusters!{ya}2:{ya}{last}<>"accept")'
         f'*(Clusters!{ya}2:{ya}{last}<>"prune")*(Clusters!{yt}2:{yt}{last}=""))'),
    ]
    wsum.append(["Check", "Value"])
    for label, formula in rows:
        wsum.append([label, formula])
    wsum.column_dimensions["A"].width = 44
    wsum.column_dimensions["B"].width = 14
    for row in wsum.iter_rows():
        for cell in row:
            cell.font = font
    wsum["A1"].font = wsum["B1"].font = bold
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


# ------------------------------------------------------------------------------------ read back
def read_workbook(path: str | Path) -> dict[str, list[dict[str, Any]]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=False)
    out = {}
    for name in ("Clusters", "Types", "Seeds"):
        ws = wb[name]
        header = [c.value for c in ws[1]]
        out[name] = [dict(zip(header, r)) for r in ws.iter_rows(min_row=2, values_only=True)
                     if r and r[0] not in (None, "")]
    return out


def schema_from_review(review: dict[str, list[dict[str, Any]]], claims: dict[str, dict[str, Any]], *,
                       genre: str) -> dict[str, Any]:
    """The reviewer's decisions -> schema dict. Raises ValueError listing every undecided or inconsistent row."""
    problems: list[str] = []
    types = {str(t["Type id"]).strip(): t for t in review["Types"]}
    sources: dict[str, list[str]] = {t: [] for t in types}
    pruned = []
    for row in review["Clusters"]:
        cid, action = row["Cluster"], (row.get("Your action") or "").strip().lower()
        if not action:
            problems.append(f"cluster {cid} undecided")
            continue
        if action not in ACTIONS:
            problems.append(f"cluster {cid}: unknown action {action!r}")
            continue
        if action == "accept":
            action, targets = row["Proposed action"], _split_list(row["Proposed type"])
            if action == "review":
                problems.append(f"cluster {cid}: 'accept' on a cluster with no proposal")
                continue
        else:
            targets = _split_list(row.get("Your type"))
        if action == "prune":
            pruned.append({"cluster": cid, "size": row["Size"], "reason": row.get("Notes") or row.get("Reason") or ""})
            continue
        if not targets:
            problems.append(f"cluster {cid}: action {action} needs a type")
        for t in targets:
            if t not in types:
                problems.append(f"cluster {cid}: type {t!r} is not on the Types sheet")
            else:
                sources[t].append(cid)
    seed_mapping: dict[str, dict[str, Any]] = {}
    for row in review["Seeds"]:
        sid, mapping = row["Seed"], (str(row.get("Your mapping") or "")).strip()
        if not mapping:
            problems.append(f"seed {sid} undecided")
            continue
        if mapping.lower() == "accept":
            if row.get("Proposed exclusion"):
                seed_mapping[sid] = {"exclude": row["Proposed exclusion"]}
                continue
            mapping = row.get("Proposed types") or ""
            if not mapping:
                problems.append(f"seed {sid}: 'accept' but nothing was proposed")
                continue
        if mapping.lower().startswith("exclude"):
            reason = mapping.split(":", 1)[1].strip() if ":" in mapping else ""
            seed_mapping[sid] = {"exclude": reason or (row.get("Notes") or "")}
        else:
            ts = _split_list(mapping)
            problems += [f"seed {sid}: type {t!r} is not on the Types sheet" for t in ts if t not in types]
            seed_mapping[sid] = {"types": ts}
    if problems:
        raise ValueError(f"{len(problems)} problem(s) in the review, e.g. undecided rows:\n  " + "\n  ".join(problems))
    out_types = []
    for tid, t in types.items():
        positives = [str(t.get(k) or "").strip() for k in ("Positive example 1", "Positive example 2")]
        out_types.append({
            "id": tid, "name": str(t.get("Name") or tid).strip(), "definition": str(t.get("Definition") or "").strip(),
            "positive_examples": [p for p in positives if p], "near_miss": str(t.get("Near miss") or "").strip(),
            "severity_weight": float(t.get("Severity weight") or 1.0),
            "section_hint": str(t.get("Section hint") or "").strip(),
            "label_source": _label_list(t.get("Label source")), "source_clusters": sources[tid]})
    return {"genre": genre, "types": out_types, "seed_mapping": seed_mapping, "pruned_clusters": pruned}


# ------------------------------------------------------------------------------------ gate checks
def validate_schema(schema: dict[str, Any], *, min_types: int, max_types: int,
                    require_label_source: bool) -> list[str]:
    """Every reason the schema fails the §3.6 gate (coverage is checked separately, task 3.5). Empty = passes."""
    problems = []
    types = schema.get("types", [])
    if not min_types <= len(types) <= max_types:
        problems.append(f"{len(types)} types; the gate needs {min_types}-{max_types}")
    seen = set()
    for t in types:
        tid = t.get("id")
        if tid in seen:
            problems.append(f"duplicate type id {tid}")
        seen.add(tid)
        for field in REQUIRED:
            if not str(t.get(field) or "").strip():
                problems.append(f"{tid}: {field} missing")
        if len([p for p in t.get("positive_examples", []) if str(p).strip()]) < 2:
            problems.append(f"{tid}: 2 positive examples needed")
        if not isinstance(t.get("severity_weight"), (int, float)):
            problems.append(f"{tid}: severity_weight missing")
        if require_label_source and not t.get("label_source"):
            problems.append(f"{tid}: label_source missing (OPP-115 values, or 'teacher')")
    for sid, m in (schema.get("seed_mapping") or {}).items():
        if not m.get("types") and not str(m.get("exclude") or "").strip():
            problems.append(f"seed {sid}: neither mapped to a type nor excluded with a reason")
    return problems


def _place(position: float) -> str:
    return ("near the start" if position < 0.25 else "in the first half" if position < 0.5
            else "in the second half" if position < 0.75 else "near the end")


def complete_proposals(proposals: dict[str, Any], clusters: list[dict[str, Any]], names: dict[str, dict[str, Any]],
                       centroids: "np.ndarray") -> dict[str, Any]:
    """Fill a proposed type's missing positive examples (the teacher's representative claims of its clusters, one
    per cluster first), near-miss (the representative claim of the nearest cluster that belongs to another type)
    and section hint (the clusters' most common heading and mean position). Filled fields are listed under `auto`
    so the workbook can say which values a human wrote."""
    import copy

    import numpy as np

    out = copy.deepcopy(proposals)
    ids = [c["id"] for c in clusters]
    by_id = {c["id"]: c for c in clusters}
    nn = nearest_clusters(centroids, ids, k=len(ids) - 1) if len(ids) > 1 else {}
    owner = {c: t["id"] for t in out.get("types", []) for c in t.get("clusters", [])}
    for t in out.get("types", []):
        src = [c for c in t.get("clusters", []) if c in by_id]
        if not src:
            continue
        auto = []
        if not t.get("positive_examples"):
            reps = [[r["claim_id"] for r in names.get(c, {}).get("representative_claims", [])] for c in src]
            picks = [r[0] for r in reps if r] + [x for r in reps for x in r[1:]]
            t["positive_examples"] = list(dict.fromkeys(picks))[:2]
            auto.append("positive_examples")
        if not t.get("near_miss"):
            for other, _ in nn.get(src[0], []):
                if owner.get(other) != t["id"] and names.get(other, {}).get("representative_claims"):
                    t["near_miss"] = names[other]["representative_claims"][0]["claim_id"]
                    break
            auto.append("near_miss")
        if not t.get("section_hint"):
            heads = {}
            for c in src:
                for h, n in by_id[c]["heading_top"]:
                    if h != "(no heading)":
                        heads[h] = heads.get(h, 0) + n
            pos = float(np.average([by_id[c]["position_mean"] for c in src], weights=[by_id[c]["size"] for c in src]))
            top = max(heads, key=heads.get) if heads else None
            if top and heads[top] < 0.1 * sum(by_id[c]["size"] for c in src):
                top = None  # a heading used by few of the claims (one paper's own heading) is not a placement hint
            t["section_hint"] = (f"under a heading like '{top}', " if top else "") + _place(pos) + " of the document"
            auto.append("section_hint")
        if auto:
            t["auto"] = auto
    return out
